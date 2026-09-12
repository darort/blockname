# helper_server.py
# Requires: pip install flask openpyxl pillow pywin32 reportlab
# Note: pywin32 (Windows COM) is optional — not required for REQUEST REFUND.xlsx

from flask import Flask, request, jsonify, send_file
import os, base64, io, re, datetime, traceback, sys, subprocess, threading, json, shlex
from openpyxl import load_workbook, Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage

# Optional Windows COM for Excel -> PDF (kept for convert_xlsx_to_pdf)
try:
    import win32com.client as win32
    import pythoncom
    WIN32_AVAILABLE = True
except Exception:
    WIN32_AVAILABLE = False

APP = Flask(__name__)

OUT_FOLDER = os.path.abspath("generated_pdfs")
os.makedirs(OUT_FOLDER, exist_ok=True)

TEMPLATE_PAYMENT = "Return Payment.xlsx"      # must be in same folder
TEMPLATE_REPRINT = "REQUEST REFUND.xlsx"     # updated to .xlsx per your change

# Do not force downscale — preserve original pixels by default.
# If you want to re-enable shrinking, set this to an integer pixel size (e.g. 1024).
MAX_IMG_PX = None

# helper config file: stores installed start_helper.bat path (text)
HELPER_CONFIG = os.path.join(OUT_FOLDER, "start_helper_path.txt")

# Process handle and lock
_helper_proc = None
_helper_proc_lock = threading.Lock()

# allow only the three signs removal for Return Payment
REMOVE_SIGNS_RE = re.compile(r'[@✓✕]')

def _clean_for_payment(s):
    if not s: return s
    return REMOVE_SIGNS_RE.sub('', str(s)).strip()

def _save_photo_from_base64(b64, out_path):
    """
    Save base64 -> file and try to embed high-DPI metadata (300 DPI).
    Returns (width, height) in pixels (or (None,None) on error).
    """
    if not b64:
        return None, None
    if ',' in b64:
        _, b64 = b64.split(',', 1)
    data = base64.b64decode(b64)
    with open(out_path, "wb") as f:
        f.write(data)

    # try to embed high DPI / keep quality, and return image size
    return _preserve_image_quality(out_path)


def _preserve_image_quality(path):
    """
    Ensure saved image has high-DPI metadata and return (width_px, height_px).
    Does not enlarge or downscale the image — only re-saves with DPI metadata and high JPEG quality.
    """
    try:
        with PILImage.open(path) as im:
            ow, oh = im.size
            fmt = (im.format or "").upper()

            save_kwargs = {}
            # set 300 DPI metadata when possible
            save_kwargs["dpi"] = (300, 300)

            # JPEG: set high quality; PNG: use optimize flag
            if fmt in ("JPEG", "JPG"):
                save_kwargs["quality"] = 95
                save_kwargs["optimize"] = True
                # ensure RGB mode for JPEG
                if im.mode in ("RGBA", "LA"):
                    im = im.convert("RGB")
            elif fmt == "PNG":
                save_kwargs["optimize"] = True
            else:
                # other formats: still try to write DPI
                pass

            try:
                im.save(path, **save_kwargs)
            except Exception:
                # if save with kwargs fails, try saving with only DPI
                try:
                    im.save(path, dpi=(300,300))
                except Exception:
                    pass

            return ow, oh
    except Exception:
        # fallback: try to open again to read size, else return None
        try:
            with PILImage.open(path) as im:
                return im.size
        except Exception:
            return None, None


def _resize_keep_aspect(path, max_px=MAX_IMG_PX):
    """
    Avoid downscaling unless max_px is an int > 0 and image is larger than max_px.
    Always (re)save image with high-DPI metadata and high JPEG quality to preserve visual sharpness.
    Returns (width_px, height_px).
    """
    try:
        with PILImage.open(path) as im:
            ow, oh = im.size
            fmt = (im.format or "").upper()

            # Only shrink if an explicit numeric max_px is provided and image is larger.
            if isinstance(max_px, (int, float)) and max_px > 0 and (ow > max_px or oh > max_px):
                scale = min(max_px/ow, max_px/oh, 1.0)
                nw, nh = int(ow*scale), int(oh*scale)
                if (nw, nh) != (ow, oh):
                    im = im.resize((nw, nh), PILImage.LANCZOS)
                    try:
                        im.save(path, dpi=(300,300))
                    except Exception:
                        im.save(path)
                    return nw, nh

            # otherwise: preserve original pixel dimensions but ensure high-quality save + DPI metadata
            save_kwargs = {}
            save_kwargs["dpi"] = (300, 300)

            if fmt in ("JPEG", "JPG"):
                save_kwargs["quality"] = 95
                save_kwargs["optimize"] = True
                if im.mode in ("RGBA", "LA"):
                    im = im.convert("RGB")
            elif fmt == "PNG":
                save_kwargs["optimize"] = True

            try:
                im.save(path, **save_kwargs)
            except Exception:
                try:
                    im.save(path, dpi=(300,300))
                except Exception:
                    pass

            return ow, oh
    except Exception:
        try:
            with PILImage.open(path) as im:
                return im.size
        except Exception:
            return None, None


def _remove_images_at_anchor(ws, anchor):
    # remove images with matching anchor attribute (best-effort)
    kept = []
    for img in getattr(ws, "_images", []):
        try:
            a = getattr(img, "anchor", None)
            if a == anchor:
                # skip (delete)
                continue
        except Exception:
            pass
        kept.append(img)
    ws._images = kept

def write_return_payment_xlsx(data, photo_bytes, out_basepath):
    """
    Writes Return Payment Excel using TEMPLATE_PAYMENT.
    out_basepath is full path without extension. returns final xlsx path.
    """
    if not os.path.exists(TEMPLATE_PAYMENT):
        raise FileNotFoundError(f"Template {TEMPLATE_PAYMENT} not found in folder.")

    wb = load_workbook(TEMPLATE_PAYMENT)
    ws = wb.active

    # cell mapping (same as you had)
    cell_map = {
        "name": "B4",
        "gender": "C4",
        "dob": "D4",
        "passport": "E4",
        "nationality": "F4",
        "firstVisa": "G4",
        "lastVisa": "H4",
        "lineCode": "L4",
        "price": "I4",
        "recall": "J4",
        "penal": "K4",
        "total": "I6"
    }

    # clean/assign values
    for k, cell in cell_map.items():
        if k == "total":
            try:
                total = float(data.get("price", 0) or 0) + float(data.get("recall", 0) or 0) + float(data.get("penal", 0) or 0)
            except Exception:
                total = data.get("total","")
            ws[cell] = total
        else:
            val = data.get(k, "")
            # clean only for payment fields (names/dates possibly containing special signs)
            if isinstance(val, str):
                val = _clean_for_payment(val)
            ws[cell] = val

    # photo insertion
    if photo_bytes:
        photo_path = out_basepath + "_photo.png"
        with open(photo_path, "wb") as f:
            f.write(photo_bytes)
        _resize_keep_aspect(photo_path, MAX_IMG_PX)
        # remove old images anchored at B8 (best-effort)
        _remove_images_at_anchor(ws, "B8")
        img = XLImage(photo_path)

        # Best-effort: set display width to the same visual width as before (360px)
        # while the embedded file keeps full pixel detail. If you want a different visible width,
        # change display_width_px accordingly.
        try:
            with PILImage.open(photo_path) as _im:
                ow, oh = _im.size
            display_width_px = 360
            img.width = display_width_px
            img.height = int(display_width_px * (oh / ow)) if ow else img.height
        except Exception:
            # if anything fails, continue without forcing display dims
            pass

        img.anchor = "B8"
        ws.add_image(img)

    out_xlsx = out_basepath + ".xlsx"
    wb.save(out_xlsx)
    return out_xlsx

def write_return_reprint_xlsx(name, line, amount, photo_bytes, out_basepath):
    """
    Uses TEMPLATE_REPRINT (REQUEST REFUND.xlsx), writes fields to C5/D5/F5/F6,
    inserts photo near B8, and returns path to generated .xlsx.

    This implementation assumes TEMPLATE_REPRINT is a valid .xlsx file (as you've converted it).
    """
    if not os.path.exists(TEMPLATE_REPRINT):
        # helpful error to debug quickly
        raise FileNotFoundError(f"Template {TEMPLATE_REPRINT} not found in folder. Please place it next to helper_server.py")

    # Load the template workbook (preserves formatting/layout)
    wb = load_workbook(TEMPLATE_REPRINT)
    ws = wb.active

    # Write values into the cells used by your template
    ws["C5"] = name
    ws["D5"] = line
    ws["F5"] = amount
    ws["F6"] = amount

    # image insertion if provided, anchor at B8
    if photo_bytes:
        photo_path = out_basepath + "_photo.png"
        with open(photo_path, "wb") as f:
            f.write(photo_bytes)
        _resize_keep_aspect(photo_path, MAX_IMG_PX)
        _remove_images_at_anchor(ws, "B8")
        img = XLImage(photo_path)

        # Best-effort: keep visible width same as previous setup (360px) while embedding high-res image.
        try:
            with PILImage.open(photo_path) as _im:
                ow, oh = _im.size
            display_width_px = 360
            img.width = display_width_px
            img.height = int(display_width_px * (oh / ow)) if ow else img.height
        except Exception:
            pass

        img.anchor = "B8"
        ws.add_image(img)

    out_xlsx = out_basepath + ".xlsx"
    wb.save(out_xlsx)
    return out_xlsx

def convert_xlsx_to_pdf(excel_xlsx_path):
    """
    Convert using Excel COM on Windows if available. Returns pdf path.
    On failure, create a trivial PDF fallback and return it.
    """
    pdf_path = excel_xlsx_path.replace(".xlsx", ".pdf")
    if WIN32_AVAILABLE and sys.platform.startswith("win"):
        try:
            pythoncom.CoInitialize()
            excel = win32.DispatchEx("Excel.Application")
            excel.Visible = False
            wb = excel.Workbooks.Open(excel_xlsx_path)
            # 0 = PDF
            wb.ExportAsFixedFormat(0, pdf_path)
            wb.Close(False)
            excel.Quit()
            pythoncom.CoUninitialize()
            return pdf_path
        except Exception as e:
            print("COM conversion failed:", e)
            traceback.print_exc()
            try:
                pythoncom.CoUninitialize()
            except Exception:
                pass

    # Fallback: make a simple PDF (ReportLab-like minimal)
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        from reportlab.lib.units import mm
        c = canvas.Canvas(pdf_path, pagesize=A4)
        w, h = A4
        c.setFont("Helvetica", 10)
        c.drawString(20*mm, h - 20*mm, f"Generated PDF fallback for: {os.path.basename(excel_xlsx_path)}")
        c.showPage()
        c.save()
        return pdf_path
    except Exception as e:
        print("Fallback PDF generation failed:", e)
        traceback.print_exc()
        return None

# -------------------- New helper management endpoints --------------------
def _read_helper_path():
    if os.path.exists(HELPER_CONFIG):
        try:
            with open(HELPER_CONFIG, "r", encoding="utf-8") as f:
                p = f.read().strip()
            if p:
                return p
        except Exception:
            pass
    # fallback default (if a previously uploaded file is kept)
    default_path = os.path.join(OUT_FOLDER, "start_helper.bat")
    if os.path.exists(default_path):
        return default_path
    return None

def _write_helper_path(path):
    with open(HELPER_CONFIG, "w", encoding="utf-8") as f:
        f.write(path)

def _is_helper_running():
    global _helper_proc
    if _helper_proc is None:
        return False
    return _helper_proc.poll() is None

def _start_helper_process(path):
    """
    Start the batch in silent/no-window mode. Returns Popen object (or raises).
    """
    global _helper_proc
    # windows: hide console
    creationflags = 0
    devnull = subprocess.DEVNULL
    if sys.platform.startswith("win"):
        CREATE_NO_WINDOW = 0x08000000
        creationflags = CREATE_NO_WINDOW
        # use shell=True to run .bat
        _helper_proc = subprocess.Popen(path, shell=True, stdout=devnull, stderr=devnull, creationflags=creationflags)
    else:
        # Unix-like
        _helper_proc = subprocess.Popen([path], stdout=devnull, stderr=devnull)
    return _helper_proc

def _stop_helper_process():
    """
    Try to stop the process (and its tree) cleanly. Returns True on success.
    """
    global _helper_proc
    if _helper_proc is None:
        return False
    pid = _helper_proc.pid
    try:
        if sys.platform.startswith("win"):
            # taskkill to ensure tree killed
            subprocess.run(["taskkill", "/F", "/T", "/PID", str(pid)], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        else:
            # Unix
            _helper_proc.terminate()
            _helper_proc.wait(timeout=3)
    except Exception:
        try:
            _helper_proc.kill()
        except Exception:
            pass
    finally:
        _helper_proc = None
    return True

@APP.route("/install_start_helper", methods=["POST"])
def install_start_helper():
    """
    Accepts multipart/form-data with file='file' (the start_helper.bat).
    Saves it into OUT_FOLDER and records path in HELPER_CONFIG.
    Returns {"path": saved_path}
    """
    try:
        if "file" not in request.files:
            return jsonify({"error": "file required"}), 400
        f = request.files["file"]
        filename = f.filename or "start_helper.bat"
        # ensure safe filename - keep only basename
        filename = os.path.basename(filename)
        saved_path = os.path.join(OUT_FOLDER, filename)
        f.save(saved_path)
        # try to make executable bit (unix)
        try:
            os.chmod(saved_path, 0o755)
        except Exception:
            pass
        _write_helper_path(saved_path)
        return jsonify({"path": saved_path})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/set_helper_path", methods=["POST"])
def set_helper_path():
    """
    Accepts JSON {"path": "<absolute-path-to-start_helper.bat>"}.
    Verifies the file exists and writes it to HELPER_CONFIG.
    Returns {"path": path} on success.
    """
    try:
        payload = request.get_json(force=True, silent=True) or {}
        path = payload.get("path")
        if not path:
            return jsonify({"error": "path required"}), 400
        # normalize
        path = os.path.abspath(path)
        if not os.path.exists(path):
            return jsonify({"error": "path not found", "path": path}), 400
        # optional: ensure it's a .bat (not necessary but helpful)
        if not path.lower().endswith(".bat"):
            return jsonify({"error": "path must point to a .bat file", "path": path}), 400

        _write_helper_path(path)
        return jsonify({"path": path})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/start_helper", methods=["POST"])
def start_helper():
    """
    Start the installed helper batch. Returns status and pid.
    """
    try:
        path = _read_helper_path()
        if not path or not os.path.exists(path):
            return jsonify({"error": "start_helper.bat not installed on server. Upload via /install_start_helper or set via /set_helper_path"}), 400

        with _helper_proc_lock:
            if _is_helper_running():
                return jsonify({"running": True, "pid": _helper_proc.pid, "path": path})
            p = _start_helper_process(path)
            return jsonify({"started": True, "pid": p.pid, "path": path})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/stop_helper", methods=["POST"])
def stop_helper():
    """
    Stop the running helper batch (if any).
    """
    try:
        with _helper_proc_lock:
            if not _is_helper_running():
                return jsonify({"stopped": True, "running": False})
            ok = _stop_helper_process()
            return jsonify({"stopped": ok, "running": False})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/helper_status", methods=["GET"])
def helper_status():
    """
    Return whether helper is running and which path is installed.
    """
    try:
        path = _read_helper_path()
        running = _is_helper_running()
        pid = _helper_proc.pid if running else None
        return jsonify({"running": running, "pid": pid, "path": path})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/open_folder", methods=["POST"])
def open_folder():
    """
    Request body JSON: { "path": "<file path>" }  or { "folder": "<folder path>" }
    The server will try to open the containing folder (Windows Explorer / macOS Finder / xdg-open).
    """
    try:
        payload = request.get_json(force=True, silent=True) or {}
        path = payload.get("path")
        folder = payload.get("folder")
        if path:
            if not os.path.exists(path):
                return jsonify({"error": "path not found", "path": path}), 400
            folder = os.path.dirname(path)
        if not folder:
            return jsonify({"error": "folder not specified"}), 400
        if not os.path.exists(folder):
            return jsonify({"error": "folder not found", "folder": folder}), 400

        if sys.platform.startswith("win"):
            os.startfile(folder)
        elif sys.platform.startswith("darwin"):
            subprocess.Popen(["open", folder])
        else:
            # assume linux with xdg-open
            subprocess.Popen(["xdg-open", folder])

        return jsonify({"opened": True, "folder": folder})
    except Exception as e:
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

# -------------------- End helper management endpoints --------------------

# CORS - allow extension to talk to us
@APP.after_request
def add_cors(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type,Authorization"
    response.headers["Access-Control-Allow-Methods"] = "GET,POST,OPTIONS"
    return response

@APP.route("/", methods=["GET"])
def health():
    return jsonify({"status":"ok"})

@APP.route("/return_payment", methods=["POST","OPTIONS"])
def return_payment():
    """
    Accepts:
      - JSON with {"data": {...}, "photo": "data:image/png;base64,..."}  (from popup.js JSON flow)
      - or multipart/form-data with 'file' (image) and form fields (if extension uses file upload)
    Returns json with xlsx_path and pdf_path (full paths).
    Also attempts to open the folder of the created file automatically (if possible).
    """
    try:
        # determine request type
        photo_bytes = None
        data = {}

        if request.is_json:
            payload = request.get_json()
            data = payload.get("data", {})
            photo_b64 = payload.get("photo")
            if photo_b64:
                if "," in photo_b64:
                    _, photo_b64 = photo_b64.split(",", 1)
                photo_bytes = base64.b64decode(photo_b64)
        else:
            # multipart/form-data: fields in form, file in files
            # gather data fields that might be present
            for k in ["name","gender","dob","passport","nationality","firstVisa","lastVisa","lineCode","price","recall","penal","total"]:
                if k in request.form:
                    data[k] = request.form.get(k)
            if "file" in request.files:
                f = request.files["file"]
                photo_bytes = f.read()

        # prepare out base name: Line (Name) Date
        line = _clean_for_payment(data.get("lineCode") or data.get("line") or "LINE")
        name = _clean_for_payment(data.get("name") or "worker")
        today = datetime.datetime.now().strftime("%d-%b-%Y")
        out_base = f"{line} ({name}) {today}"
        out_basepath = os.path.join(OUT_FOLDER, out_base)

        xlsx_path = write_return_payment_xlsx(data, photo_bytes, out_basepath)
        pdf_path = convert_xlsx_to_pdf(xlsx_path)

        # try to open folder of created file (best-effort)
        try:
            folder = os.path.dirname(pdf_path or xlsx_path)
            if folder and os.path.exists(folder):
                if sys.platform.startswith("win"):
                    os.startfile(folder)
                elif sys.platform.startswith("darwin"):
                    subprocess.Popen(["open", folder])
                else:
                    subprocess.Popen(["xdg-open", folder])
        except Exception:
            pass

        return jsonify({"xlsx_path": xlsx_path, "pdf_path": pdf_path})
    except Exception as e:
        print("ERROR /return_payment:", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

@APP.route("/return_reprint", methods=["POST","OPTIONS"])
def return_reprint():
    """
    Accepts:
      - JSON with { "name":..., "line":..., "amount":..., "photo": "data:...base64" }
      - or multipart form with 'name','line','amount' and file 'file'
    Will attempt to open the folder of the generated file.
    """
    try:
        photo_bytes = None
        name = None
        line = None
        amount = None

        if request.is_json:
            payload = request.get_json()
            name = payload.get("name")
            line = payload.get("line")
            amount = payload.get("amount")
            photo_b64 = payload.get("photo")
            if photo_b64:
                if "," in photo_b64:
                    _, photo_b64 = photo_b64.split(",", 1)
                photo_bytes = base64.b64decode(photo_b64)
        else:
            # form data
            name = request.form.get("name")
            line = request.form.get("line")
            amount = request.form.get("amount")
            if "file" in request.files:
                f = request.files["file"]
                photo_bytes = f.read()

        if not name or not line:
            return jsonify({"error":"name and line required"}), 400

        today = datetime.datetime.now().strftime("%d-%b-%Y")
        out_base = f"REPRINT ({name}) {line} ({today})"
        out_basepath = os.path.join(OUT_FOLDER, out_base)

        xlsx_path = write_return_reprint_xlsx(name, line, amount, photo_bytes, out_basepath)
        pdf_path = convert_xlsx_to_pdf(xlsx_path)

        # try to open folder of created file (best-effort)
        try:
            folder = os.path.dirname(pdf_path or xlsx_path)
            if folder and os.path.exists(folder):
                if sys.platform.startswith("win"):
                    os.startfile(folder)
                elif sys.platform.startswith("darwin"):
                    subprocess.Popen(["open", folder])
                else:
                    subprocess.Popen(["xdg-open", folder])
        except Exception:
            pass

        return jsonify({"xlsx_path": xlsx_path, "pdf_path": pdf_path})
    except Exception as e:
        print("ERROR /return_reprint:", e)
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500

if __name__ == "__main__":
    print("Helper running on http://127.0.0.1:5000")
    APP.run(host="127.0.0.1", port=5000)
